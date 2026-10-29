"""
Style Atlas - ComfyUI node that builds prompts from an XLSX style library.

The node works on its own (Python only). The optional JS file adds linked
dropdowns and live counts; if it fails, the node still loads and runs.
"""
import os
import re
import glob
import random

NODE_DIR = os.path.dirname(os.path.abspath(__file__))
LIBRARY_NAME = "style_atlas_library.xlsx"

MODIFIER_CATEGORIES = [
    "Quality", "Mood & Atmosphere", "Composition & Shot", "Genre",
    "Lighting", "Palette", "Architectural Style", "Vehicle Design", "Product Design",
]
# Subject-specific modifiers: only offered when the matching subject is chosen
MODIFIER_SUBJECT = {
    "Architectural Style": "Architecture",
    "Vehicle Design": "Automotive",
    "Product Design": "Product",
}

RANDOM = "🎲 Random"
NONE = "None"
ALL = "All"

POSITIONS = {
    "prompt, style": "{prompt}, {style}",
    "style, prompt": "{style}, {prompt}",
    "Style: style. Subject: prompt": "Style: {style} Subject: {prompt}",
    "Subject: prompt. Style: style": "Subject: {prompt}. Style: {style}",
    "prompt. style": "{prompt}. {style}",
    "prompt, in the style of style": "{prompt}, in the style of {style}",
    "prompt + new line + style": "{prompt}\n{style}",
    "style only": "{style}",
    "custom template": None,
}

VERSIONS = ["Short (fallback Long)", "Long (fallback Short)", "Short only", "Long only"]
REALISM = [ALL, "1", "2", "3", "4", "5"]
REALISM_MODE = ["Exact", "±1"]

COUNT_RE = re.compile(r"\s*\(\d+\)\s*$")


def clean(value):
    """Remove the ' (123)' count suffix added by the JS for display."""
    return COUNT_RE.sub("", str(value or "")).strip()


# --------------------------------------------------------------------------
# Library loading (cached, reloads automatically when the xlsx is modified)
# --------------------------------------------------------------------------
_cache = {"path": None, "mtime": None, "rows": [], "error": None}


def library_path():
    p = os.path.join(NODE_DIR, LIBRARY_NAME)
    if os.path.isfile(p):
        return p
    found = sorted(f for f in glob.glob(os.path.join(NODE_DIR, "*.xlsx"))
                   if not os.path.basename(f).startswith("~$"))
    return found[0] if found else None


def load_library():
    path = library_path()
    if not path:
        _cache.update(rows=[], error=f"No library found. Put {LIBRARY_NAME} in {NODE_DIR}")
        return _cache["rows"]
    mtime = os.path.getmtime(path)
    if _cache["path"] == path and _cache["mtime"] == mtime:
        return _cache["rows"]
    try:
        import openpyxl  # imported here so a missing package never blocks the node
    except ImportError:
        _cache.update(rows=[], error="openpyxl is not installed (pip install openpyxl)")
        return _cache["rows"]
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb["Styles"] if "Styles" in wb.sheetnames else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it)]

        def col(*names):
            for n in names:
                if n in header:
                    return header.index(n)
            return None

        c = dict(name=col("Style Name"), cat=col("Category"), sub=col("Subcategory"),
                 era=col("Era"), real=col("Realism (1-5)", "Realism"),
                 short=col("Short Version"), long=col("Long Version"),
                 subject=col("Subject"), layer=col("Layer"))

        def get(r, key):
            i = c[key]
            return r[i] if i is not None and i < len(r) else None

        rows = []
        for r in it:
            name = get(r, "name")
            if not name:
                continue
            cat = str(get(r, "cat") or "").strip()
            short = str(get(r, "short") or "").strip()
            long_ = str(get(r, "long") or "").strip()
            if not short and not long_:
                continue
            try:
                real = int(get(r, "real"))
            except (TypeError, ValueError):
                real = 3
            layer = str(get(r, "layer") or "").strip() or (
                "Modifier" if cat in MODIFIER_CATEGORIES else "Base Style")
            rows.append(dict(
                name=str(name).strip(), cat=cat, sub=str(get(r, "sub") or "").strip(),
                era=str(get(r, "era") or "Any").strip(), real=real,
                short=short, long=long_,
                subject=str(get(r, "subject") or "Main").strip(), layer=layer))
        wb.close()
        _cache.update(path=path, mtime=mtime, rows=rows, error=None)
    except Exception as e:  # never crash ComfyUI because of the spreadsheet
        _cache.update(path=None, mtime=None, rows=[], error=f"Could not read {os.path.basename(path)}: {e}")
    return _cache["rows"]


def public_data():
    """Light data sent to the JS (no prompt text)."""
    rows = load_library()
    return {
        "error": _cache["error"],
        "modifier_categories": MODIFIER_CATEGORIES,
        "modifier_subject": MODIFIER_SUBJECT,
        "rows": [dict(n=r["name"], c=r["cat"], s=r["sub"], j=r["subject"], e=r["era"],
                      r=r["real"], m=1 if r["layer"] == "Modifier" else 0,
                      hs=1 if r["short"] else 0, hl=1 if r["long"] else 0) for r in rows],
    }


# --------------------------------------------------------------------------
# Filtering helpers (same logic as the JS, so results match the counts)
# --------------------------------------------------------------------------
def version_ok(r, version):
    if version == "Short only":
        return bool(r["short"])
    if version == "Long only":
        return bool(r["long"])
    return True


def pick_text(r, version):
    if version.startswith("Long"):
        return r["long"] or ("" if version == "Long only" else r["short"])
    return r["short"] or ("" if version == "Short only" else r["long"])


def subject_ok(r, subject):
    return subject == ALL or r["subject"] == subject


def realism_ok(r, realism, mode):
    if realism == ALL:
        return True
    target = int(realism)
    return r["real"] == target if mode == "Exact" else abs(r["real"] - target) <= 1


def strip_name(text):
    """'Style Name: description.' -> 'description.'"""
    head, sep, tail = text.partition(":")
    return tail.strip() if sep and len(head) <= 80 else text


def finish(text):
    text = text.strip()
    return text if not text or text[-1] in ".!?" else text + "."


# --------------------------------------------------------------------------
# Node
# --------------------------------------------------------------------------
class StyleAtlasNode:
    @classmethod
    def INPUT_TYPES(cls):
        rows = load_library()
        base = [r for r in rows if r["layer"] != "Modifier"]
        mods = [r for r in rows if r["layer"] == "Modifier"]
        subjects = [ALL] + sorted({r["subject"] for r in rows if r["subject"]},
                                  key=lambda s: (s != "Main", s))
        cats = [ALL] + list(dict.fromkeys(r["cat"] for r in base))
        subs = [ALL] + sorted({r["sub"] for r in base if r["sub"]})
        eras = [ALL] + sorted({r["era"] for r in base if r["era"]})
        styles = [RANDOM] + sorted(r["name"] for r in base)

        req = {
            "prompt": ("STRING", {"multiline": True, "default": ""}),
            "seed": ("INT", {"default": 0, "min": 0, "max": 0xFFFFFFFFFFFFFFFF}),
            "version": (VERSIONS,),
            "subject": (subjects,),
            "realism": (REALISM,),
            "realism_mode": (REALISM_MODE,),
            "era": (eras,),
            "category": (cats,),
            "subcategory": (subs,),
            "style": (styles,),
            "style_position": (list(POSITIONS.keys()),),
            "include_style_names": ("BOOLEAN", {"default": True}),
        }
        for mc in MODIFIER_CATEGORIES:
            names = sorted(r["name"] for r in mods if r["cat"] == mc)
            req[cls.mod_key(mc)] = ([NONE, RANDOM] + names,)
        return {
            "required": req,
            "optional": {
                "custom_template": ("STRING", {
                    "multiline": False,
                    "default": "{prompt}, {style}",
                    "tooltip": "Used when style_position = custom template. Placeholders: {prompt} {style} {modifiers}",
                }),
            },
        }

    RETURN_TYPES = ("STRING", "STRING", "STRING")
    RETURN_NAMES = ("prompt", "style_text", "info")
    FUNCTION = "build"
    CATEGORY = "Style Atlas/Prompt"

    @staticmethod
    def mod_key(category):
        return "mod_" + re.sub(r"[^a-z0-9]+", "_", category.lower()).strip("_")

    @classmethod
    def VALIDATE_INPUTS(cls, **kwargs):
        # Dropdown labels carry live counts from the JS ("Painting (272)"),
        # so the built-in list check is skipped; values are cleaned in build().
        return True

    def build(self, prompt, seed, version, subject, realism, realism_mode, era,
              category, subcategory, style, style_position, include_style_names,
              custom_template="{prompt}, {style}", **mods):
        rows = load_library()
        if not rows:
            msg = _cache["error"] or "Style library is empty"
            return (prompt, "", f"⚠ {msg}")

        subject, realism, era = clean(subject), clean(realism), clean(era)
        category, subcategory, style = clean(category), clean(subcategory), clean(style)
        rng = random.Random(seed)

        # ---- base style -----------------------------------------------------
        pool = [r for r in rows if r["layer"] != "Modifier"
                and subject_ok(r, subject)
                and realism_ok(r, realism, realism_mode)
                and (era == ALL or r["era"] == era)
                and (category == ALL or r["cat"] == category)
                and (subcategory == ALL or r["sub"] == subcategory)
                and version_ok(r, version)]
        chosen = None
        if style and style != RANDOM:
            chosen = next((r for r in rows if r["name"] == style), None)
        elif pool:
            chosen = rng.choice(pool)

        info = []
        parts = []
        if chosen:
            text = pick_text(chosen, version)
            parts.append(finish(text if include_style_names else strip_name(text)))
            info.append(f"Style: {chosen['name']} [{chosen['cat']} / {chosen['sub']}] "
                        f"(pool {len(pool)})")
        else:
            info.append("Style: none matched the current filters")

        # ---- modifiers ----------------------------------------------------
        mod_parts = []
        for mc in MODIFIER_CATEGORIES:
            val = clean(mods.get(self.mod_key(mc), NONE))
            if not val or val == NONE:
                continue
            need = MODIFIER_SUBJECT.get(mc)
            if need and subject not in (ALL, need):
                continue
            cands = [r for r in rows if r["cat"] == mc and version_ok(r, version)
                     and (r["subject"] in ("Main", subject) or subject == ALL)]
            if val == RANDOM:
                m = rng.choice(cands) if cands else None
            else:
                m = next((r for r in rows if r["cat"] == mc and r["name"] == val), None)
            if m:
                t = pick_text(m, version)
                mod_parts.append(finish(t if include_style_names else strip_name(t)))
                info.append(f"{mc}: {m['name']}")

        modifiers = " ".join(mod_parts)
        style_text = " ".join(p for p in [*parts, modifiers] if p)

        # ---- assemble -------------------------------------------------------
        prompt = (prompt or "").strip()
        if style_position == "custom template":
            tpl = custom_template or "{prompt}, {style}"
            if "{modifiers}" in tpl:
                out = tpl.replace("{prompt}", prompt).replace("{style}", " ".join(parts)) \
                         .replace("{modifiers}", modifiers)
            else:
                out = tpl.replace("{prompt}", prompt).replace("{style}", style_text)
        else:
            tpl = POSITIONS.get(clean(style_position), "{prompt}, {style}")
            if not prompt:
                out = style_text
            elif not style_text:
                out = prompt
            else:
                p = prompt.rstrip(" .,") if "{prompt}," in tpl or "{prompt}." in tpl else prompt
                out = tpl.replace("{prompt}", p).replace("{style}", style_text)
        out = re.sub(r"[ \t]+", " ", out).replace(" ,", ",").replace("..", ".").strip()
        return (out, style_text, "\n".join(info))


# --------------------------------------------------------------------------
# API route for the JS (registered only if the ComfyUI server is available)
# --------------------------------------------------------------------------
try:
    from server import PromptServer
    from aiohttp import web

    @PromptServer.instance.routes.get("/style_atlas/data")
    async def _style_atlas_data(request):
        return web.json_response(public_data())
except Exception:
    pass
